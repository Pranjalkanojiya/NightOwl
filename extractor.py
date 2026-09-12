import os
import fitz
from pptx import Presentation


def extract_pdf(path):
    pages = []

    document = fitz.open(path)

    for i, page in enumerate(document):
        text = page.get_text("text").strip()

        pages.append({
            "text": text,
            "source": os.path.basename(path),
            "page": i + 1,
            "type": "pdf"
        })

    document.close()

    return pages


def extract_pptx(path):
    presentation = Presentation(path)
    pages = []

    for i, slide in enumerate(presentation.slides):
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_parts.append(shape.text)

        pages.append({
            "text": "\n".join(text_parts).strip(),
            "source": os.path.basename(path),
            "page": i + 1,
            "type": "pptx"
        })

    return pages


def extract_text(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()

    return [{
        "text": text,
        "source": os.path.basename(path),
        "page": 1,
        "type": "text"
    }]


def extract(path):
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return extract_pdf(path)

    if ext == ".pptx":
        return extract_pptx(path)

    if ext in [".txt", ".md"]:
        return extract_text(path)

    return []